# -*- coding: utf-8 -*-
# 生成《炼钢北区周边系统清单及通讯方式.xlsx》
import sys, os, io

XLSX_SKILL_DIR = r'C:\Users\AIGC\.zcode\cli\plugins\cache\zcode-plugins-official\document-skills\0.1.4\skills\xlsx'
for sub in [XLSX_SKILL_DIR, os.path.join(XLSX_SKILL_DIR, 'templates')]:
    if sub not in sys.path:
        sys.path.insert(0, sub)

from base import *  # noqa
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

OUT = r'D:\work\company\太钢二炼钢\analysis\炼钢北区周边系统清单及通讯方式.xlsx'
wb = Workbook()
wb.properties.creator = "Z.ai"

# ================= Sheet1 周边系统总表 =================
ws = wb.active
ws.title = '周边系统总表'
H1 = ['序号', '对端系统', '厂商', '节点', '现有接口', '改造后接口', '电文解析结构/格式', '代码侧印证']
rows1 = [
    (1, '太钢制造管理系统', '宝信', '0R', 'PO', 'xcom直接通讯', 'JSON(GBK)', 'cm_002021 接收制造命令(PSSM);框架分发数据增加北区'),
    (2, '物流运输管理系统', '宝信', '31', 'PO', '测试:xcom直接通讯\n运行:xcom+IXBUS', 'JSON(GBK),子母表形式', 'f_xxpa04_snd(铁运车皮,XXPA04)、f_xxsm01_snd(发货码单,7000S4)'),
    (3, '资源综合利用系统', '宝信', '33', 'PO', 'xcom直接通讯', 'JSON(GBK),TLD协议', 'mmsm2a_snd/n(给资源铁区发消耗)'),
    (4, '铁区动态管控系统', '宝信', '32', 'PO', 'xcom直接通讯', 'JSON(GBK),TLD协议', 'cm_32t809_rcv(铁区库存)、21B004/21B005/21B006 发送'),
    (5, '能源动态管控系统', '宝信/宝能', 'F0', '前置机+IXBUS', 'xcom+IXBUS', '—', 'cm_nyme01/02_rcv(能源消耗接收)、t8f004/005/007/fsb/se 能源动态发送'),
    (6, '计量系统', '宝信自动化', 'EW', 'PO', 'restful(不走xcom)', 'JSON', 'cm_ewt801_rcv、f_mmsm81_d021_handle_rcv、cm_b02103_rcv;已部署'),
    (7, '智慧质量系统', '宝信', '23', '前置机+IXBUS', 'xcom直接通讯', '默认格式', 'QMTS 工序制造标准 11 工序双向(002113~002123)、cm_200001_rcv'),
    (8, '炼钢二厂南区MES', '中冶赛迪', 'T7', 'PO', 'xcom+南区前置机', 'TGPO/29S', '南区接收:北区MES->前置机->PO->南区MES;南区发送反向;北/南标识32处'),
    (9, 'Broner排程系统', '—', '—', '产销->PO->MES(转写Broner本地表)->Broner读取;发送反向经MES->L2', '不直接通讯', '—', '代码库未见Broner直连(与官方结论一致)'),
    (10, '炼钢一厂不锈钢MES', '山宝', 'U0', 'PO', 'xcom+一钢前置机', 'TGPO/29S', '路由路径同南区MES'),
    (11, '炼钢一厂碳钢MES', '山宝', 'V0', 'PO', 'xcom+一钢前置机', 'TGPO/29S', '同上'),
    (12, '热连轧2250MES', 'Psi', 'P3', 'PO', 'xcom直接通讯', 'TGPO/29S', '北区->2250:框架直接写2250接口表;2250->北区:写xcom的sendbuff表'),
    (13, '热连轧1549MES', '宝信', 'M7', 'PO', 'xcom直接通讯', 'JSON(GBK),TLD协议', '—'),
    (14, '型材厂MES', '宝信', 'P1', 'PO', 'xcom直接通讯', 'JSON(GBK),TLD协议', '—'),
    (15, '不锈线材厂MES', '宝信', 'P0', 'PO', 'xcom直接通讯', 'JSON(GBK),TLD协议', '—'),
    (16, '4300中厚板厂MES', '宝信', '24', '前置机+IXBUS', 'xcom直接通讯', 'JSON(GBK),TLD协议', '—'),
    (17, '太钢中板MES', '宝信', 'M9', '前置机+IXBUS', 'xcom直接通讯', 'JSON(GBK),TLD协议', '—'),
    (18, 'L2 com系统(含板坯库L2)', '普瑞特/中冶南方/宝信', 'E2', 'DBLink', 'xcom+独立前置机', 'SHIGANG', 'cm_1a1013/15/21转炉实绩收、cm_1a1087连铸切断收、f_mmsm3301板坯原始数据收;发T8E2S1计划状态/PAS1P1出钢计划/PAS1P2铸坯命令/PAS1P3钢种变更/t8e2原料系列'),
    (19, '检化验L2系统前置通讯机', '—', 'EJ', 'DBLink', 'xcom+独立前置机', '—', '检化验采用Oracle,通过ID号=5判断是否给MES;QMTS检查结果接收(210003)'),
    (20, '炼钢二厂iPlat集控平台', '宝信自动化', '16', '数据抽取', '数据抽取', '—', '官方表注:需要协调'),
    (21, '倒罐站数采系统', '—', 'ED', '新建', '新建', '默认', '北区MES->前置机receviedbuff;倒罐站->写前置sentbuff表'),
    (22, '不锈冷轧系统', '—', '—', 'DBLink', 'DBLink(保留)', '—', '原10.11.132.212需转换为10.162/163/164(新库网段);现有北区10.13.133.136'),
    (23, 'TPS工艺模型(代码侧补充)', '宝信', '—', '—', 'REST/JSON', '—', 'f_epex_call_rest_tpsmodel*(APBD 3个),CRestClient(conn,"MZX_TEST"),服务CP_MODEL_TG,端点在iPlat4C EX04'),
    (24, '宝武chat消息推送(代码侧补充)', '宝信', '—', 'EPEX电文', 'EPEX电文', '含wxurl字段', 'QMTS f_push_baowu_chat/_dd/_pro、qm_push_baowu_wb,4个程序,消息通知通道'),
]
setup_sheet(ws, title='炼钢北区(二炼钢)周边系统总表 — 2026-09-08 梳理(官方接口整理表20230705+代码证据)', last_col=1 + len(H1))
for c, h in enumerate(H1, 2):
    ws.cell(row=4, column=c, value=h)
style_header_row(ws, row_num=4, col_start=2, col_end=1 + len(H1))
for i, r in enumerate(rows1):
    for c, v in enumerate(r, 2):
        ws.cell(row=5 + i, column=c, value=v)
    style_data_row(ws, row_num=5 + i, col_start=2, col_end=1 + len(H1), row_index=i)
ws.freeze_panes = 'C5'
ws.page_setup.orientation = 'landscape'
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.print_title_rows = '4:4'
auto_fit_columns(ws, min_width=8, max_width=40, header_row=4, data_start_row=5)
auto_fit_row_heights(ws, header_row=4, data_start_row=5)

# ================= Sheet2 通道机制说明 =================
ws2 = wb.create_sheet('通道机制说明')
H2 = ['通道', '机制', '现状/说明', '状态']
rows2 = [
    ('PO 电文中间件', '宝信平台电文路由节点,各系统现有主通道', '电文经 PO 路由,改造后逐步被 xcom 直替', '存量主通道'),
    ('XCOM 文件传输', '对端/前置机写 sendbuff 表(发送)、读 receviedbuff 表(接收);北区接收程序以 REC_CREATOR="XCOM" 落库', '改造后主通道;计量系统已按 restful 先行部署', '目标形态'),
    ('前置机', '独立前置机(L2 com/检化验)、南区前置机、一钢前置机;热连轧2250为框架直接写接口表', '按对端保留', '在用'),
    ('IXBUS', '宝信集成总线', '能源/4300中厚板/太钢中板/智慧质量现有在用', '存量在用'),
    ('RESTful(HTTP/JSON)', 'iPlat4C CRestClient(conn,"MZX_TEST"),服务名 CP_MODEL_TG;端点配置在平台 EX04 外部系统注册', '计量系统、TPS 工艺模型(APBD 3个程序)在用', '在用'),
    ('DBLink(Oracle)', '数据库级链接,应用 SQL 中无表@链名写法', 'L2 com(E2)、检化验(EJ)待改 xcom+独立前置机;不锈冷轧保留并切新网段', '存量改造点'),
    ('宝武chat推送', 'EPEX 电文形态(含 wxurl 字段),QMTS 4个推送程序', '内部告警/通知', '在用'),
    ('— 未使用通道 —', 'Oracle 表@链名、database link 建链语句、MQ、FTP、Socket/TCP 直连、UNC 共享目录、MSSQL linked server、客户端 WebService', '全库检索均无证据', '未使用'),
]
setup_sheet(ws2, title='通讯通道机制说明', last_col=1 + len(H2))
for c, h in enumerate(H2, 2):
    ws2.cell(row=4, column=c, value=h)
style_header_row(ws2, row_num=4, col_start=2, col_end=1 + len(H2))
for i, r in enumerate(rows2):
    for c, v in enumerate(r, 2):
        ws2.cell(row=5 + i, column=c, value=v)
    style_data_row(ws2, row_num=5 + i, col_start=2, col_end=1 + len(H2), row_index=i)
ws2.freeze_panes = 'C5'
ws2.page_setup.orientation = 'landscape'
ws2.page_setup.fitToWidth = 1
ws2.page_setup.fitToHeight = 0
auto_fit_columns(ws2, min_width=8, max_width=40, header_row=4, data_start_row=5)
auto_fit_row_heights(ws2, header_row=4, data_start_row=5)

# ================= Sheet3 电文程序统计 =================
ws3 = wb.create_sheet('电文程序统计')
H3 = ['模块', '电文程序总数', '发送(snd)', '接收(rcv)']
tpath = os.path.join(os.path.dirname(OUT), 'dameng-migration', 'batches', 'comm_snd_rcv_inventory.tsv')
inv = [l.rstrip('\n').split('\t') for l in io.open(tpath, encoding='utf-8')][1:]
mods = {}
for m, fn, d, desc, tc, e in inv:
    a = mods.setdefault(m, [0, 0, 0])
    a[0] += 1
    if d == 'snd':
        a[1] += 1
    else:
        a[2] += 1
rows3 = sorted(((m, a[0], a[1], a[2]) for m, a in mods.items()), key=lambda x: -x[1])
setup_sheet(ws3, title='电文收发程序统计(EPEX/BM2 框架,214 个以 BM2F_ENTERACE_TELE 注册为电文入口)', last_col=1 + len(H3))
for c, h in enumerate(H3, 2):
    ws3.cell(row=4, column=c, value=h)
style_header_row(ws3, row_num=4, col_start=2, col_end=1 + len(H3))
for i, r in enumerate(rows3):
    for c, v in enumerate(r, 2):
        cell = ws3.cell(row=5 + i, column=c, value=v)
    style_data_row(ws3, row_num=5 + i, col_start=2, col_end=1 + len(H3), row_index=i)
    for c in range(3, 6):
        ws3.cell(row=5 + i, column=c).alignment = align_number()
tr = 5 + len(rows3)
ws3.cell(row=tr, column=2, value='合计')
for c in range(3, 6):
    L = get_column_letter(c)
    ws3.cell(row=tr, column=c, value=f'=SUM({L}5:{L}{tr-1})')
style_total_row(ws3, row_num=tr, col_start=2, col_end=1 + len(H3))
ws3.freeze_panes = 'C5'
auto_fit_columns(ws3, min_width=8, max_width=30, header_row=4, data_start_row=5)
auto_fit_row_heights(ws3, header_row=4, data_start_row=5)

# ================= Sheet4 电文程序明细 =================
ws4 = wb.create_sheet('电文程序明细')
H4 = ['模块', '程序文件', '方向', '功能描述', '电文号', '电文入口']
setup_sheet(ws4, title='电文收发程序明细(444 个,静态盘点)', last_col=1 + len(H4))
for c, h in enumerate(H4, 2):
    ws4.cell(row=4, column=c, value=h)
style_header_row(ws4, row_num=4, col_start=2, col_end=1 + len(H4))
for i, r in enumerate(inv):
    for c, v in enumerate(r, 2):
        ws4.cell(row=5 + i, column=c, value=v)
    style_data_row(ws4, row_num=5 + i, col_start=2, col_end=1 + len(H4), row_index=i)
ws4.freeze_panes = 'C5'
ws4.page_setup.orientation = 'landscape'
ws4.page_setup.fitToWidth = 1
ws4.page_setup.fitToHeight = 0
ws4.print_title_rows = '4:4'
auto_fit_columns(ws4, min_width=8, max_width=40, header_row=4, data_start_row=5)
auto_fit_row_heights(ws4, header_row=4, data_start_row=5)

# ================= Sheet5 DM8迁移关系与待确认事项 =================
ws5 = wb.create_sheet('DM8迁移与待确认')
H5 = ['类别', '事项', '说明']
rows5 = [
    ('DM8迁移关系', '接口改造方案即官方表"改造后接口"列', '代码侧 444 个电文程序的 SQL 已全量入账并转换(台账已转换 384 行)'),
    ('DM8迁移关系', 'DBLink 3 个对端为关键路径', 'L2 com、检化验改 xcom+独立前置机(应用侧无感);不锈冷轧保留 DBLink,链目标切至新网段 10.162/163/164'),
    ('DM8迁移关系', 'XCOM/前置机/IXBUS/PO 与数据库方言无关', '文件与消息通道,无 SQL 方言问题'),
    ('DM8迁移关系', 'RESTful 与宝武chat 与数据库无关', '走 HTTP/电文通知'),
    ('待确认', '平台侧配置', 'iPlat4C EX04 外部系统注册、PO 电文路由、XCOM 作业与节点定义均在部署侧,代码库内无配置'),
    ('待确认', '网络与防火墙', '节点间连通性按改造后网段核对(不锈冷轧 10.162/163/164;现有北区 10.13.133.136)'),
    ('待确认', '行车系统衔接', 'WMSM 天车命令程序群 27 个(cranecmd 系列)为库表级管理,与行车系统物理衔接无电文/Socket 证据,疑经共享库表轮询'),
    ('待确认', '南区路由', '与南区MES(T7)/一钢(U0/V0)的前置机转发配置'),
    ('说明', '分析方法', '官方《外部接口接口整理_20230705.xlsx》+ 源码静态证据互证;未连接生产环境核实;基于 2026-09-08 工作区快照'),
]
setup_sheet(ws5, title='DM8 迁移关系与待确认事项', last_col=1 + len(H5))
for c, h in enumerate(H5, 2):
    ws5.cell(row=4, column=c, value=h)
style_header_row(ws5, row_num=4, col_start=2, col_end=1 + len(H5))
for i, r in enumerate(rows5):
    for c, v in enumerate(r, 2):
        ws5.cell(row=5 + i, column=c, value=v)
    style_data_row(ws5, row_num=5 + i, col_start=2, col_end=1 + len(H5), row_index=i)
ws5.freeze_panes = 'C5'
ws5.page_setup.orientation = 'landscape'
ws5.page_setup.fitToWidth = 1
ws5.page_setup.fitToHeight = 0
auto_fit_columns(ws5, min_width=8, max_width=40, header_row=4, data_start_row=5)
auto_fit_row_heights(ws5, header_row=4, data_start_row=5)

wb.save(OUT)
print('saved:', OUT)
